from collections import Counter
from pathlib import Path

from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.core.models import Ciudadano, Manzana


class Command(BaseCommand):
    help = (
        "Importa el padrón inicial de ciudadanos desde un archivo Excel. "
        "Usa --dry-run para validar sin modificar la base de datos."
    )

    EXPECTED_HEADERS = {
        "nombre",
        "apellido_paterno",
        "apellido_materno",
        "sexo",
        "codigo_mzn",
    }

    SEXO_IMPORTACION = "NO_ESPECIFICADO"

    def add_arguments(self, parser):
        parser.add_argument(
            "archivo",
            type=str,
            help="Ruta al archivo .xlsx que contiene la hoja 'Importar'.",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Valida todo el archivo sin modificar la base de datos.",
        )

    def handle(self, *args, **options):
        archivo = Path(options["archivo"])
        dry_run = options["dry_run"]

        if not archivo.exists():
            raise CommandError(
                f"No existe el archivo: {archivo}"
            )

        if archivo.suffix.lower() != ".xlsx":
            raise CommandError(
                "El archivo debe tener extensión .xlsx"
            )

        try:
            workbook = load_workbook(
                filename=archivo,
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise CommandError(
                f"No se pudo abrir el archivo Excel: {exc}"
            ) from exc

        if "Importar" not in workbook.sheetnames:
            raise CommandError(
                "El archivo no contiene la hoja requerida "
                "'Importar'."
            )

        worksheet = workbook["Importar"]

        # ---------------------------------------------------------
        # 1. Encabezados
        # ---------------------------------------------------------

        try:
            headers = [
                self._clean_header(cell.value)
                for cell in next(
                    worksheet.iter_rows(
                        min_row=1,
                        max_row=1,
                    )
                )
            ]
        except StopIteration as exc:
            raise CommandError(
                "La hoja 'Importar' está vacía."
            ) from exc

        missing_headers = (
            self.EXPECTED_HEADERS - set(headers)
        )

        if missing_headers:
            raise CommandError(
                "Faltan columnas requeridas en la hoja "
                "'Importar': "
                + ", ".join(sorted(missing_headers))
            )

        header_indexes = {
            header: index
            for index, header in enumerate(headers)
        }

        # ---------------------------------------------------------
        # 2. Leer filas y validar datos básicos
        # ---------------------------------------------------------

        rows = []
        errors = []
        duplicate_keys = Counter()
        manzana_cache = {}

        for excel_row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):
            if self._row_is_empty(values):
                continue

            row_data = {
                header: (
                    values[index]
                    if index < len(values)
                    else None
                )
                for header, index in header_indexes.items()
            }

            try:
                citizen_data = self._build_citizen_data(
                    row_data=row_data,
                    manzana_cache=manzana_cache,
                )
            except ValueError as exc:
                errors.append(
                    f"Fila {excel_row_number}: {exc}"
                )
                continue

            duplicate_key = self._duplicate_key(
                citizen_data
            )

            duplicate_keys[duplicate_key] += 1

            rows.append(
                {
                    "excel_row": excel_row_number,
                    "data": citizen_data,
                    "duplicate_key": duplicate_key,
                }
            )

        # ---------------------------------------------------------
        # 3. Detectar duplicados
        # ---------------------------------------------------------

        duplicates = {
            key
            for key, count in duplicate_keys.items()
            if count > 1
        }

        duplicate_rows = [
            row
            for row in rows
            if row["duplicate_key"] in duplicates
        ]

        # ---------------------------------------------------------
        # 4. Validación Django
        #
        # Los duplicados se excluyen de esta primera importación.
        # Los demás registros continúan normalmente.
        # ---------------------------------------------------------

        validation_errors = []
        valid_rows = []

        for row in rows:

            if row["duplicate_key"] in duplicates:
                continue

            citizen = Ciudadano(
                **row["data"]
            )

            try:
                citizen.full_clean()

            except ValidationError as exc:
                validation_errors.append(
                    self._format_validation_error(
                        row["excel_row"],
                        exc,
                    )
                )
                continue

            valid_rows.append(citizen)

        # ---------------------------------------------------------
        # 5. Resumen
        # ---------------------------------------------------------

        total_rows = (
            len(rows) + len(errors)
        )

        duplicate_count = len(
            duplicate_rows
        )

        self.stdout.write("")
        self.stdout.write(
            self.style.MIGRATE_HEADING(
                "=== VALIDACIÓN DEL PADRÓN ==="
            )
        )

        self.stdout.write(
            f"Archivo: {archivo}"
        )

        self.stdout.write(
            "Hoja: Importar"
        )

        self.stdout.write(
            f"Filas leídas: {total_rows}"
        )

        self.stdout.write(
            f"Filas válidas para importar: "
            f"{len(valid_rows)}"
        )

        self.stdout.write(
            f"Filas duplicadas excluidas: "
            f"{duplicate_count}"
        )

        self.stdout.write(
            f"Errores de datos: {len(errors)}"
        )

        self.stdout.write(
            "Errores de validación Django: "
            f"{len(validation_errors)}"
        )

        # ---------------------------------------------------------
        # 6. Mostrar duplicados como advertencia
        # ---------------------------------------------------------

        if duplicate_rows:

            self.stdout.write("")

            self.stdout.write(
                self.style.WARNING(
                    "=== DUPLICADOS EXCLUIDOS ==="
                )
            )

            self.stdout.write(
                "Estos registros NO serán importados "
                "en esta primera carga."
            )

            duplicate_groups = {}

            for row in duplicate_rows:

                duplicate_groups.setdefault(
                    row["duplicate_key"],
                    [],
                ).append(
                    row["excel_row"]
                )

            for key, row_numbers in duplicate_groups.items():

                self.stdout.write(
                    f"Filas "
                    f"{', '.join(map(str, row_numbers))}: "
                    f"{key}"
                )

        # ---------------------------------------------------------
        # 7. Mostrar errores de datos
        # ---------------------------------------------------------

        if errors:

            self.stdout.write("")

            self.stdout.write(
                self.style.ERROR(
                    "=== ERRORES DE DATOS ==="
                )
            )

            for error in errors:
                self.stdout.write(error)

        # ---------------------------------------------------------
        # 8. Mostrar errores Django
        # ---------------------------------------------------------

        if validation_errors:

            self.stdout.write("")

            self.stdout.write(
                self.style.ERROR(
                    "=== ERRORES DE VALIDACIÓN DJANGO ==="
                )
            )

            for error in validation_errors:
                self.stdout.write(error)

        # ---------------------------------------------------------
        # 9. Los duplicados NO bloquean.
        #    Los errores reales sí bloquean.
        # ---------------------------------------------------------

        has_errors = bool(
            errors
            or validation_errors
        )

        if has_errors:

            self.stdout.write("")

            self.stdout.write(
                self.style.ERROR(
                    "Estado: NO SE REALIZÓ LA IMPORTACIÓN."
                )
            )

            self.stdout.write(
                "Corrige los errores indicados y "
                "vuelve a ejecutar --dry-run."
            )

            return

        # ---------------------------------------------------------
        # 10. Dry run
        # ---------------------------------------------------------

        if dry_run:

            self.stdout.write("")

            self.stdout.write(
                self.style.SUCCESS(
                    "Estado: VALIDACIÓN EXITOSA. "
                    "No se modificó la base de datos."
                )
            )

            self.stdout.write(
                f"Se importarían {len(valid_rows)} "
                "ciudadanos."
            )

            if duplicate_count:
                self.stdout.write(
                    self.style.WARNING(
                        f"Se excluirían {duplicate_count} "
                        "filas por duplicidad."
                    )
                )

            self.stdout.write(
                "Para realizar la importación real, "
                "ejecuta el comando sin --dry-run."
            )

            return

        # ---------------------------------------------------------
        # 11. Importación real
        # ---------------------------------------------------------

        try:

            with transaction.atomic():

                Ciudadano.objects.bulk_create(
                    valid_rows,
                    batch_size=500,
                )

        except Exception as exc:

            raise CommandError(
                "La importación fue cancelada y no se "
                f"guardaron registros: {exc}"
            ) from exc

        # ---------------------------------------------------------
        # 12. Resultado final
        # ---------------------------------------------------------

        self.stdout.write("")

        self.stdout.write(
            self.style.SUCCESS(
                "=== IMPORTACIÓN COMPLETADA ==="
            )
        )

        self.stdout.write(
            f"Ciudadanos creados: {len(valid_rows)}"
        )

        self.stdout.write(
            f"Filas excluidas por duplicidad: "
            f"{duplicate_count}"
        )

        self.stdout.write(
            self.style.SUCCESS(
                "La base de datos fue actualizada "
                "correctamente."
            )
        )

    # =============================================================
    # Métodos auxiliares
    # =============================================================

    @staticmethod
    def _clean_header(value):
        if value is None:
            return ""

        return str(value).strip()

    @staticmethod
    def _clean_text(value):
        if value is None:
            return ""

        return " ".join(
            str(value).strip().split()
        )

    @staticmethod
    def _row_is_empty(values):
        return all(
            value is None
            or str(value).strip() == ""
            for value in values
        )

    def _required_text(
        self,
        value,
        field_name,
    ):
        cleaned = self._clean_text(value)

        if not cleaned:
            raise ValueError(
                f"el campo {field_name} está vacío."
            )

        return cleaned

    def _build_citizen_data(
        self,
        row_data,
        manzana_cache,
    ):
        nombre = self._required_text(
            row_data.get("nombre"),
            "nombre",
        )

        apellido_paterno = self._required_text(
            row_data.get("apellido_paterno"),
            "apellido_paterno",
        )

        apellido_materno = self._clean_text(
            row_data.get("apellido_materno")
        )

        codigo_mzn = self._required_text(
            row_data.get("codigo_mzn"),
            "codigo_mzn",
        )

        sexo = self._clean_text(
            row_data.get("sexo")
        )

        # El archivo ya viene normalizado.
        if sexo != self.SEXO_IMPORTACION:
            raise ValueError(
                f"sexo debe ser "
                f"'{self.SEXO_IMPORTACION}', "
                f"pero se encontró "
                f"'{sexo or '(vacío)'}'."
            )

        # Buscar Manzana por su clave.
        if codigo_mzn not in manzana_cache:

            try:
                manzana_cache[codigo_mzn] = (
                    Manzana.objects.get(
                        clave=codigo_mzn
                    )
                )

            except Manzana.DoesNotExist as exc:
                raise ValueError(
                    f"Codigo MZN '{codigo_mzn}' "
                    "no existe en Manzana.clave."
                ) from exc

            except Manzana.MultipleObjectsReturned as exc:
                raise ValueError(
                    f"Codigo MZN '{codigo_mzn}' "
                    "tiene más de una Manzana."
                ) from exc

        return {
            "nombre": nombre,
            "apellido_paterno": apellido_paterno,
            "apellido_materno": apellido_materno,
            "sexo": self.SEXO_IMPORTACION,
            "manzana": manzana_cache[codigo_mzn],
        }

    @staticmethod
    def _duplicate_key(data):
        """
        Clave para detectar duplicados dentro del Excel:

        nombre + apellido_paterno + apellido_materno
        """

        return " ".join(
            part
            for part in (
                data["nombre"],
                data["apellido_paterno"],
                data["apellido_materno"],
            )
            if part
        ).upper()

    @staticmethod
    def _format_validation_error(
        row_number,
        error,
    ):
        messages = []

        for field, field_errors in (
            error.message_dict.items()
        ):

            for message in field_errors:
                messages.append(
                    f"{field}: {message}"
                )

        if not messages:
            messages.append(
                str(error)
            )

        return (
            f"Fila {row_number}: "
            + " | ".join(messages)
        )